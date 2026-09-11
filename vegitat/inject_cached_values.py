#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Überträgt die von LibreOffice berechneten Formelwerte (aus einer durchgerechneten
Kopie) als zwischengespeicherte Werte in die Original-Datei von openpyxl.

Warum: openpyxl schreibt Formeln ohne Ergebniswerte. Excel rechnet beim Öffnen
selbst, aber Datei-Vorschauen (Handy, Cloud, Mail) zeigen sonst leere Zellen.
Die Originaldatei behält dabei ihre komplette Formatierung, Diagramme, Tabellen
und Dropdowns exakt so, wie openpyxl sie geschrieben hat.

Aufruf:  python3 inject_cached_values.py <original.xlsx> <durchgerechnet.xlsx> <ziel.xlsx>
"""
import datetime as dt
import re
import sys
import zipfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

src, calc, dst = sys.argv[1:4]
values = load_workbook(calc, data_only=True)

# Blattname -> sheetN.xml (Reihenfolge in workbook.xml == Reihenfolge der rId in den rels)
with zipfile.ZipFile(src) as z:
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
sheet_names = re.findall(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="([^"]+)"', wbxml)
rel_map = dict(re.findall(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels))
rel_map.update({k: v for v, k in re.findall(r'<Relationship [^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
part_for = {}
for name, rid in sheet_names:
    target = rel_map[rid]
    part_for["xl/" + target.lstrip("/").replace("xl/", "", 1) if not target.startswith("/xl/") else target[1:]] = name.replace("&amp;", "&")

cell_re = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)
stats = {"num": 0, "str": 0, "bool": 0, "empty": 0}


def render(m, ws):
    ref, attrs, formula = m.group(1), m.group(2), m.group(3)
    v = ws[ref].value
    attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
    if v is None:
        stats["empty"] += 1
        return f'<c r="{ref}"{attrs} t="str"><f>{formula}</f><v></v></c>'
    if isinstance(v, bool):
        stats["bool"] += 1
        return f'<c r="{ref}"{attrs} t="b"><f>{formula}</f><v>{int(v)}</v></c>'
    if isinstance(v, (dt.datetime, dt.date)):
        v = to_excel(v)
    if isinstance(v, (int, float)):
        stats["num"] += 1
        return f'<c r="{ref}"{attrs}><f>{formula}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
    stats["str"] += 1
    return f'<c r="{ref}"{attrs} t="str"><f>{formula}</f><v>{escape(str(v))}</v></c>'


with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in part_for:
            ws = values[part_for[item.filename]]
            xml = data.decode("utf-8")
            xml = cell_re.sub(lambda m: render(m, ws), xml)
            data = xml.encode("utf-8")
        zout.writestr(item, data)
print("geschrieben:", dst, stats)
