#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Baut die Excel-Datei "Vegitat_Finanzuebersicht.xlsx":
Einnahmen-/Ausgaben-Cockpit für die Vegitat-Filialen Zürich und Luzern (CHF).

Aufruf:
    python3 build_finanzuebersicht.py [Zielpfad] [--test]

--test setzt die Filial-Auswahl auf den Blättern Monats-/Jahresübersicht auf
einzelne Filialen, damit verify_finanzuebersicht.py beide Formelzweige prüfen kann.
"""
import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

# ----------------------------------------------------------------- Konstanten
FONT = "Arial"
FIRST = 5            # erste Datenzeile im Journal "Buchungen"
LAST = 5000          # letzte Journalzeile, die von den Formeln erfasst wird
N_FIL, N_INC, N_EXP, N_PAY = 10, 10, 25, 10
EXAMPLE_YEAR = 2025

FILIALEN = ["Zürich", "Luzern"]
EINNAHMEN = [
    "Verkauf Laden – Bar",
    "Verkauf Laden – Karte/Twint",
    "Lieferdienste (Uber Eats, Just Eat, Smood)",
    "Online-Bestellungen (Webshop)",
    "Catering & Events",
    "Gutscheine",
    "Sonstige Einnahmen",
]
AUSGABEN = [
    "Wareneinkauf Lebensmittel",
    "Wareneinkauf Getränke",
    "Verpackung & Takeaway-Material",
    "Löhne & Gehälter",
    "Sozialversicherungen (AHV/IV/EO, ALV, BVG, UVG, FAK)",
    "Miete Ladenlokal",
    "Nebenkosten (Strom, Wasser, Heizung)",
    "Internet, Telefon & Software (Kasse)",
    "Marketing & Werbung",
    "Lieferplattform-Gebühren (Kommissionen)",
    "Karten- & Twint-Gebühren",
    "Reinigung & Hygiene",
    "Reparaturen & Unterhalt",
    "Geräte & Inventar (Anschaffungen)",
    "Versicherungen",
    "Steuern & Abgaben (MWST, Gebühren)",
    "Buchhaltung & Beratung (Treuhänder)",
    "Fahrzeug & Transport",
    "Bankgebühren & Zinsen",
    "Sonstige Ausgaben",
]
ZAHLUNGSARTEN = ["Bar", "Karte (EC/Kredit)", "Twint", "Banküberweisung", "Lieferplattform", "Sonstiges"]
MONATE = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]

# Arbeitgeberbeiträge (Schätzwerte, siehe Hinweis im Blatt Mitarbeiter)
AG_SAETZE = [
    ("AHV/IV/EO – Arbeitgeberanteil", 0.053),
    ("ALV – Arbeitgeberanteil", 0.011),
    ("BVG / Pensionskasse – Arbeitgeberanteil (Schätzung)", 0.050),
    ("UVG (Berufsunfall) + KTG (Schätzung)", 0.015),
    ("FAK / Familienzulagen (kantonal, Schätzung)", 0.015),
]

# Farben
C_HDR, C_SEC, C_TOT, C_IN, C_TILE, C_WARN, C_NOTE = "1F6F43", "DDEBDD", "EFEFEF", "FFF2CC", "F3F7F4", "FCE4D6", "F7F7F7"
COL_BLUE, COL_ORANGE, COL_AQUA, COL_VIOLET = "2A78D6", "EB6834", "1BAF7A", "4A3AA7"

# Zahlenformate
NF_CHF = '#,##0.00;[Red]-#,##0.00;"-"'
NF_CHF0 = '"CHF "#,##0;[Red]-"CHF "#,##0;"CHF 0"'
NF_PCT = '0.0%;[Red]-0.0%;"-"'
NF_INT = '#,##0;[Red]-#,##0;"-"'
NF_DATE = "DD.MM.YYYY"
NF_MONTH = "MMM YYYY"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP = Border(top=Side(style="medium", color="404040"))


# ----------------------------------------------------------------- Helfer
def font(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def put(ws, ref, value, bold=False, size=10, color="000000", nf=None, fill_=None,
        align=None, italic=False, wrap=False, border=None):
    c = ws[ref]
    c.value = value
    c.font = font(bold, size, color, italic)
    if nf:
        c.number_format = nf
    if fill_:
        c.fill = fill(fill_)
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        c.border = border
    return c


def header_row(ws, row, col_from, labels, height=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col_from + i, value=lab)
        c.font = font(True, 10, "FFFFFF")
        c.fill = fill(C_HDR)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    if height:
        ws.row_dimensions[row].height = height


def input_cell(ws, ref, value, nf=None, align="center"):
    c = put(ws, ref, value, bold=True, color="0000FF", nf=nf, fill_=C_IN, align=align, border=BORDER)
    return c


def section(ws, row, text, last_col):
    for col in range(2, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.fill = fill(C_SEC)
        c.font = font(True)
    ws.cell(row=row, column=2, value=text)


def total_style(ws, row, col_from, col_to, fill_=C_TOT):
    for col in range(col_from, col_to + 1):
        c = ws.cell(row=row, column=col)
        c.font = font(True)
        c.fill = fill(fill_)
        c.border = TOP


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub=None):
    put(ws, "B1", text, bold=True, size=16, color=C_HDR)
    if sub:
        put(ws, "B2", sub, italic=True, color="595959")
    ws.row_dimensions[1].height = 24


# Bereiche im Journal
def R(col):
    return f"Buchungen!${col}${FIRST}:${col}${LAST}"


D_, F_, T_, K_, B_, Z_ = R("A"), R("B"), R("C"), R("D"), R("F"), R("G")


def sumifs(*crits):
    return "SUMIFS(" + ",".join([B_] + [f"{r},{c}" for r, c in crits]) + ")"


def countifs(*crits):
    return "COUNTIFS(" + ",".join([f'{B_},"<>"'] + [f"{r},{c}" for r, c in crits]) + ")"


def month_crit(y, m):
    return [(D_, f'">="&DATE({y},{m},1)'), (D_, f'"<"&DATE({y},{m}+1,1)')]


def year_crit(y):
    return [(D_, f'">="&DATE({y},1,1)'), (D_, f'"<"&DATE({y}+1,1,1)')]


def fil_switch(sel, crits, fn=sumifs):
    """'Alle' -> ohne Filialfilter (inkl. Buchungen ohne Filiale), sonst mit."""
    return f'IF({sel}="Alle",{fn(*crits)},{fn(*(list(crits) + [(F_, sel)]))})'


def label_link(ref):
    return f'=IF({ref}="","",{ref})'


# ----------------------------------------------------------------- Beispieldaten
def r05(x):
    return round(round(x * 20) / 20, 2)


def example_rows():
    """Beispielbuchungen für ein ganzes Jahr, beide Filialen (deterministisch)."""
    rnd = random.Random(2025)
    Y = EXAMPLE_YEAR
    base = {"Zürich": 54000.0, "Luzern": 36000.0}
    season = [0.92, 0.90, 0.98, 1.00, 1.04, 1.08, 1.10, 1.10, 1.02, 1.00, 0.96, 0.95]
    loehne = {"Zürich": 14840.0, "Luzern": 13300.0}
    miete = {"Zürich": 5800.0, "Luzern": 3400.0}
    fix = {  # Kategorie-Index (AUSGABEN) -> (ZH, LU, Tag, Zahlungsart, Beschreibung)
        6: (720, 480, 15, "Banküberweisung", "Strom/Wasser Monatsrechnung"),
        7: (240, 190, 5, "Karte (EC/Kredit)", "Internet, Telefon, Kassensoftware"),
        11: (380, 260, 20, "Banküberweisung", "Reinigungsfirma / Hygieneartikel"),
        14: (285, 215, 3, "Banküberweisung", "Betriebsversicherung Monatsprämie"),
        16: (350, 250, 8, "Banküberweisung", "Treuhänder Monatspauschale"),
        18: (45, 35, 0, "Banküberweisung", "Bankspesen"),
    }
    rows = []

    def add(date, fil, typ, kat, beschr, betrag, zart, beleg=""):
        rows.append({"Datum": date, "Filiale": fil, "Typ": typ, "Kategorie": kat,
                     "Beschreibung": beschr, "Betrag": r05(betrag), "Zahlungsart": zart,
                     "Beleg": beleg, "Bemerkung": "Beispiel"})

    for m in range(1, 13):
        last_day = (dt.date(Y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)).day
        for fil in FILIALEN:
            rev = base[fil] * season[m - 1] * rnd.uniform(0.95, 1.05)
            karte, twint, bar = rev * 0.45, rev * 0.15, rev * 0.20
            liefer, online = rev * 0.17, rev * 0.03
            pre = "Monatsumsatz (Beispiel – im Echtbetrieb pro Tag erfassen): "
            add(dt.date(Y, m, last_day), fil, "Einnahme", EINNAHMEN[1], pre + "Karte", karte, "Karte (EC/Kredit)")
            add(dt.date(Y, m, last_day), fil, "Einnahme", EINNAHMEN[1], pre + "Twint", twint, "Twint")
            add(dt.date(Y, m, last_day), fil, "Einnahme", EINNAHMEN[0], pre + "Bargeld", bar, "Bar")
            add(dt.date(Y, m, last_day), fil, "Einnahme", EINNAHMEN[2], "Auszahlung Lieferplattformen (Bruttobestellwert)", liefer, "Lieferplattform")
            add(dt.date(Y, m, last_day), fil, "Einnahme", EINNAHMEN[3], "Webshop-Bestellungen", online, "Karte (EC/Kredit)")
            if rnd.random() < 0.25:
                add(dt.date(Y, m, rnd.randint(2, 26)), fil, "Einnahme", EINNAHMEN[4], "Firmen-Catering (Apéro)", rnd.uniform(800, 2500), "Banküberweisung")
            # Ausgaben
            food = rev * 0.29
            add(dt.date(Y, m, 7), fil, "Ausgabe", AUSGABEN[0], "Lieferant Gemüse, Brot, Saucen (1. Hälfte)", food * 0.5, "Banküberweisung")
            add(dt.date(Y, m, 21), fil, "Ausgabe", AUSGABEN[0], "Lieferant Seitan/Cigköfte, Gewürze (2. Hälfte)", food * 0.5, "Banküberweisung")
            add(dt.date(Y, m, 10), fil, "Ausgabe", AUSGABEN[1], "Getränkelieferung", rev * 0.04, "Banküberweisung")
            add(dt.date(Y, m, 12), fil, "Ausgabe", AUSGABEN[2], "Takeaway-Boxen, Becher, Servietten", rev * 0.03, "Karte (EC/Kredit)")
            add(dt.date(Y, m, 25), fil, "Ausgabe", AUSGABEN[3], "Lohnzahlung Bruttolöhne", loehne[fil], "Banküberweisung")
            add(dt.date(Y, m, 28), fil, "Ausgabe", AUSGABEN[4], "AHV/ALV/BVG/UVG Arbeitgeberbeiträge", loehne[fil] * 0.144, "Banküberweisung")
            add(dt.date(Y, m, 1), fil, "Ausgabe", AUSGABEN[5], "Miete Ladenlokal", miete[fil], "Banküberweisung")
            for idx, (zh, lu, day, zart, beschr) in fix.items():
                amount = (zh if fil == "Zürich" else lu) * rnd.uniform(0.9, 1.1)
                add(dt.date(Y, m, day or last_day), fil, "Ausgabe", AUSGABEN[idx], beschr, amount, zart)
            if rnd.random() < 0.4:
                add(dt.date(Y, m, 18), fil, "Ausgabe", AUSGABEN[8], "Social-Media-Anzeigen / Flyer", rnd.uniform(300, 1500), "Karte (EC/Kredit)")
            add(dt.date(Y, m, last_day), fil, "Ausgabe", AUSGABEN[9], "Kommission Lieferplattformen (30 %)", liefer * 0.30, "Lieferplattform")
            add(dt.date(Y, m, last_day), fil, "Ausgabe", AUSGABEN[10], "Kartenterminal-/Twint-Gebühren", (karte + twint) * 0.014, "Banküberweisung")
            if rnd.random() < 0.25:
                add(dt.date(Y, m, rnd.randint(2, 26)), fil, "Ausgabe", AUSGABEN[12], "Reparatur Grill/Kühlung", rnd.uniform(150, 900), "Karte (EC/Kredit)")
            if m in (3, 6, 9, 12):
                add(dt.date(Y, m, 28), fil, "Ausgabe", AUSGABEN[15], f"MWST-Abrechnung Quartal {m // 3}", 3100 if fil == "Zürich" else 1950, "Banküberweisung")
            if fil == "Zürich":
                add(dt.date(Y, m, 14), fil, "Ausgabe", AUSGABEN[17], "Lieferfahrzeug Benzin/Parkplatz", rnd.uniform(250, 380), "Karte (EC/Kredit)")
            if rnd.random() < 0.3:
                add(dt.date(Y, m, rnd.randint(2, 26)), fil, "Ausgabe", AUSGABEN[19], "Kleinmaterial / Diverses", rnd.uniform(60, 300), "Bar")
    add(dt.date(Y, 3, 11), "Zürich", "Ausgabe", AUSGABEN[13], "Neuer Kühlschrank", 2400, "Banküberweisung", "R-2025-031")
    add(dt.date(Y, 8, 5), "Luzern", "Ausgabe", AUSGABEN[13], "Neuer Dönergrill", 1150, "Banküberweisung", "R-2025-080")
    rows.sort(key=lambda r: (r["Datum"], r["Filiale"]))
    # Zwei absichtlich fehlerhafte Beispiele (zeigen die automatische Prüfung):
    rows.append({"Datum": dt.date(Y, 5, 3), "Filiale": "Zürich", "Typ": "Ausgabe", "Kategorie": "Dekoration",
                 "Beschreibung": "Beispiel: Kategorie existiert nicht in 'Listen' → landet unter 'Nicht zugeordnet'",
                 "Betrag": 120.0, "Zahlungsart": "Bar", "Beleg": "", "Bemerkung": "Beispiel"})
    rows.append({"Datum": dt.date(Y, 6, 15), "Filiale": "", "Typ": "Ausgabe", "Kategorie": AUSGABEN[19],
                 "Beschreibung": "Beispiel: Filiale fehlt → orange markiert, im Dashboard als 'unvollständig' gezählt",
                 "Betrag": 50.0, "Zahlungsart": "Bar", "Beleg": "", "Bemerkung": "Beispiel"})
    return rows


MITARBEITER_BEISPIELE = [
    ("Beispiel: A. Yilmaz", "Zürich", "Filialleitung", dt.date(2021, 3, 1), 1.0, 5600, "Ja"),
    ("Beispiel: S. Keller", "Zürich", "Koch/Köchin", dt.date(2022, 8, 15), 1.0, 4600, "Ja"),
    ("Beispiel: M. Rossi", "Zürich", "Verkauf/Service", dt.date(2023, 5, 1), 0.7, 2940, "Nein"),
    ("Beispiel: L. Brunner", "Zürich", "Aushilfe (Stundenlohn, Ø)", dt.date(2024, 11, 1), 0.2, 850, "Nein"),
    ("Beispiel: E. Demir", "Luzern", "Filialleitung", dt.date(2020, 9, 1), 1.0, 5200, "Ja"),
    ("Beispiel: J. Huber", "Luzern", "Koch/Köchin", dt.date(2023, 2, 1), 1.0, 4400, "Ja"),
    ("Beispiel: N. Meier", "Luzern", "Verkauf/Service", dt.date(2024, 4, 1), 0.5, 2100, "Nein"),
    ("Beispiel: T. Aslan", "Luzern", "Aushilfe (Stundenlohn, Ø)", dt.date(2025, 1, 10), 0.2, 800, "Nein"),
]

# Zeilenlayout der Auswertungsblätter (Monats- und Jahresübersicht, Block 1)
INC_ROWS = list(range(9, 9 + N_INC))            # 9..18
R_INC_NZ, R_INC_TOT = 19, 20
R_EXP_HDR = 22
EXP_ROWS = list(range(23, 23 + N_EXP))          # 23..47
R_EXP_NZ, R_EXP_TOT = 48, 49
R_RES, R_CUM, R_MARGE, R_CNT = 51, 52, 53, 54

# Budget-Zeilen
BUD_INC = {i: 6 + i for i in range(N_INC)}       # 6..15
BUD_INC_TOT = 16
BUD_EXP = {j: 19 + j for j in range(N_EXP)}      # 19..43
BUD_EXP_TOT, BUD_RES, BUD_MARGE = 44, 46, 47

# Mitarbeiter-Zeilen
MA_FIRST, MA_N = 14, 20
MA_LAST = MA_FIRST + MA_N - 1                    # 33
MA_SUM = {"Zürich": 38, "Luzern": 39, "Gesamt": 40}


# ================================================================= Blätter
def build_listen(wb):
    ws = wb.create_sheet("Listen")
    title(ws, "Listen – Filialen, Kategorien, Zahlungsarten",
          "Hier anpassen: neue Einträge in die nächste freie Zeile schreiben. Dropdowns und Auswertungen folgen automatisch.")
    ws["A1"].value = None
    header_row(ws, 2, 1, ["Filialen", "Einnahme", "Ausgabe", "Zahlungsarten"])
    put(ws, "F2", "Filiale (Auswahl)", bold=True, color="FFFFFF", fill_=C_HDR, align="center", border=BORDER)
    put(ws, "H2", "Hinweis", bold=True, color="FFFFFF", fill_=C_HDR, align="center", border=BORDER)
    for i in range(N_FIL):
        input_cell(ws, f"A{3 + i}", FILIALEN[i] if i < len(FILIALEN) else None, align="left")
    for i in range(N_INC):
        input_cell(ws, f"B{3 + i}", EINNAHMEN[i] if i < len(EINNAHMEN) else None, align="left")
    for i in range(N_EXP):
        input_cell(ws, f"C{3 + i}", AUSGABEN[i] if i < len(AUSGABEN) else None, align="left")
    for i in range(N_PAY):
        input_cell(ws, f"D{3 + i}", ZAHLUNGSARTEN[i] if i < len(ZAHLUNGSARTEN) else None, align="left")
    put(ws, "F3", "Alle", border=BORDER)
    for i in range(N_FIL):
        put(ws, f"F{4 + i}", f'=IF(A{3 + i}="","",A{3 + i})', border=BORDER)
    notes = [
        "Gelbe Zellen = Eingabe. Die Spaltentitel 'Einnahme' und 'Ausgabe' nicht umbenennen (die Dropdowns im Journal suchen danach).",
        f"Kapazität: {N_FIL} Filialen, {N_INC} Einnahme-Kategorien, {N_EXP} Ausgabe-Kategorien, {N_PAY} Zahlungsarten.",
        "Kategorie umbenennen: zuerst im Blatt 'Buchungen' per Suchen/Ersetzen alle alten Einträge anpassen, sonst laufen sie unter 'Nicht zugeordnet'.",
        "Die Spalte 'Filiale (Auswahl)' wird automatisch gebildet (für die Filter in den Auswertungen).",
        "Die Filialvergleiche (Dashboard, Jahresübersicht) verwenden die ersten beiden Filialen dieser Liste.",
    ]
    for i, n in enumerate(notes):
        put(ws, f"H{3 + i}", n, wrap=True)
        ws.row_dimensions[3 + i].height = 15 * max(1, -(-len(n) // 95))
    widths(ws, {"A": 16, "B": 40, "C": 48, "D": 20, "E": 3, "F": 18, "G": 3, "H": 100})
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.freeze_panes = "A3"
    return ws


def build_buchungen(wb, rows):
    ws = wb.create_sheet("Buchungen")
    put(ws, "A1", "Buchungsjournal – alle Einnahmen und Ausgaben (CHF)", bold=True, size=16, color=C_HDR)
    put(ws, "A2", "Jede Einnahme und jede Ausgabe als eigene Zeile erfassen. Pflichtfelder: Datum, Filiale, Typ, Kategorie, Betrag. "
        "Betrag immer positiv – ob Einnahme oder Ausgabe, bestimmt die Spalte 'Typ'. "
        "Neue Buchung: einfach in die nächste leere Zeile unter der Tabelle schreiben. "
        "Beispielzeilen (Bemerkung = 'Beispiel') vor dem Echtbetrieb löschen.", italic=True, color="595959")
    ws.row_dimensions[1].height = 24
    headers = ["Datum *", "Filiale *", "Typ *", "Kategorie *", "Beschreibung", "Betrag (CHF) *", "Zahlungsart", "Beleg-Nr.", "Bemerkung"]
    header_row(ws, 4, 1, headers, height=30)
    # Beispielzeilen
    for i, r in enumerate(rows):
        row = FIRST + i
        vals = [r["Datum"], r["Filiale"] or None, r["Typ"], r["Kategorie"], r["Beschreibung"], r["Betrag"], r["Zahlungsart"], r["Beleg"] or None, r["Bemerkung"]]
        for j, v in enumerate(vals):
            ws.cell(row=row, column=j + 1, value=v)
    # Formatierung aller Journalzeilen (auch der noch leeren)
    for row in range(FIRST, LAST + 1):
        for col in range(1, 10):
            c = ws.cell(row=row, column=col)
            c.font = font()
            if col == 1:
                c.number_format = NF_DATE
                c.alignment = Alignment(horizontal="center")
            elif col == 6:
                c.number_format = NF_CHF
            elif col == 3:
                c.alignment = Alignment(horizontal="center")
    last_tab = FIRST + len(rows) - 1
    tab = Table(displayName="tblBuchungen", ref=f"A4:I{last_tab}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)

    # Datenvalidierung
    rng = f"{FIRST}:{LAST}"
    dv_date = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2040,12,31)",
                             allow_blank=True, showErrorMessage=True, showInputMessage=True,
                             errorTitle="Ungültiges Datum", error="Bitte ein Datum im Format TT.MM.JJJJ eingeben (2020–2040).",
                             promptTitle="Datum", prompt="Format TT.MM.JJJJ, z. B. 05.03.2026")
    dv_date.add(f"A{FIRST}:A{LAST}")
    dv_fil = DataValidation(type="list", formula1=f"OFFSET(Listen!$A$3,0,0,MAX(1,COUNTA(Listen!$A$3:$A${2 + N_FIL})),1)",
                            allow_blank=True, showErrorMessage=True, showInputMessage=True,
                            errorTitle="Unbekannte Filiale", error="Bitte eine Filiale aus der Liste wählen (Blatt 'Listen').",
                            promptTitle="Filiale", prompt="Aus der Liste wählen")
    dv_fil.add(f"B{FIRST}:B{LAST}")
    dv_typ = DataValidation(type="list", formula1='"Einnahme,Ausgabe"', allow_blank=True, showErrorMessage=True,
                            showInputMessage=True, errorTitle="Typ", error="Bitte 'Einnahme' oder 'Ausgabe' wählen.",
                            promptTitle="Typ", prompt="Einnahme oder Ausgabe")
    dv_typ.add(f"C{FIRST}:C{LAST}")
    dv_kat = DataValidation(
        type="list",
        formula1=(f"OFFSET(Listen!$B$3,0,MATCH($C{FIRST},Listen!$B$2:$C$2,0)-1,"
                  f"MAX(1,COUNTA(OFFSET(Listen!$B$3,0,MATCH($C{FIRST},Listen!$B$2:$C$2,0)-1,{N_EXP},1))),1)"),
        allow_blank=True, showErrorMessage=True, showInputMessage=True,
        errorTitle="Unbekannte Kategorie", error="Bitte eine Kategorie aus der Liste wählen. Neue Kategorien im Blatt 'Listen' anlegen.",
        promptTitle="Kategorie", prompt="Zuerst 'Typ' wählen, dann passende Kategorie")
    dv_kat.add(f"D{FIRST}:D{LAST}")
    dv_betr = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                             showInputMessage=True, errorTitle="Betrag", error="Betrag muss grösser als 0 sein. Einnahme/Ausgabe wird über 'Typ' bestimmt.",
                             promptTitle="Betrag in CHF", prompt="Immer positiv, z. B. 1234.50")
    dv_betr.add(f"F{FIRST}:F{LAST}")
    dv_pay = DataValidation(type="list", formula1=f"OFFSET(Listen!$D$3,0,0,MAX(1,COUNTA(Listen!$D$3:$D${2 + N_PAY})),1)",
                            allow_blank=True, showErrorMessage=True, errorTitle="Zahlungsart", error="Bitte eine Zahlungsart aus der Liste wählen.")
    dv_pay.add(f"G{FIRST}:G{LAST}")
    for dv in (dv_date, dv_fil, dv_typ, dv_kat, dv_betr, dv_pay):
        ws.add_data_validation(dv)

    # Bedingte Formatierung: unvollständige Zeilen orange, Typ farbig
    full = f"A{FIRST}:I{LAST}"
    ws.conditional_formatting.add(full, FormulaRule(
        formula=[f'AND($F{FIRST}<>"",OR($A{FIRST}="",$B{FIRST}="",$C{FIRST}="",$D{FIRST}=""))'], fill=fill(C_WARN), stopIfTrue=False))
    for col in ("C", "F"):
        r = f"{col}{FIRST}:{col}{LAST}"
        ws.conditional_formatting.add(r, FormulaRule(formula=[f'$C{FIRST}="Einnahme"'], font=Font(name=FONT, color="1E7B34", bold=(col == "F"))))
        ws.conditional_formatting.add(r, FormulaRule(formula=[f'$C{FIRST}="Ausgabe"'], font=Font(name=FONT, color="B02A2A", bold=(col == "F"))))

    widths(ws, {"A": 12, "B": 12, "C": 11, "D": 40, "E": 52, "F": 14, "G": 18, "H": 12, "I": 22})
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = "FFC000"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build_overview_block(ws, first_col, ncols, col_crit, sel_fil, fil_link_rows=True):
    """Kategorie-Zeilen + Totale. col_crit(col_idx) -> Liste der Datums-Kriterien für die Spalte."""
    cols = [get_column_letter(first_col + i) for i in range(ncols)]
    tot_col = get_column_letter(first_col + ncols)  # Spalte 'Total' (nur Monatsübersicht)
    # Labels
    section(ws, 8, "EINNAHMEN", first_col + ncols + 3)
    for i, r in enumerate(INC_ROWS):
        put(ws, f"B{r}", label_link(f"Listen!$B${3 + i}"))
    put(ws, f"B{R_INC_NZ}", "Nicht zugeordnet (ohne / unbekannte Kategorie)", italic=True, color="7F7F7F")
    put(ws, f"B{R_INC_TOT}", "TOTAL EINNAHMEN", bold=True)
    section(ws, R_EXP_HDR, "AUSGABEN", first_col + ncols + 3)
    for j, r in enumerate(EXP_ROWS):
        put(ws, f"B{r}", label_link(f"Listen!$C${3 + j}"))
    put(ws, f"B{R_EXP_NZ}", "Nicht zugeordnet (ohne / unbekannte Kategorie)", italic=True, color="7F7F7F")
    put(ws, f"B{R_EXP_TOT}", "TOTAL AUSGABEN", bold=True)
    put(ws, f"B{R_RES}", "ERGEBNIS (Einnahmen − Ausgaben)", bold=True)
    for ci, col in enumerate(cols):
        crit = col_crit(ci)
        for i, r in enumerate(INC_ROWS):
            f = fil_switch(sel_fil, [(T_, '"Einnahme"'), (K_, f"$B{r}")] + crit)
            put(ws, f"{col}{r}", f'=IF($B{r}="",0,{f})', nf=NF_CHF)
        put(ws, f"{col}{R_INC_TOT}", "=" + fil_switch(sel_fil, [(T_, '"Einnahme"')] + crit), nf=NF_CHF)
        put(ws, f"{col}{R_INC_NZ}", f"={col}{R_INC_TOT}-SUM({col}{INC_ROWS[0]}:{col}{INC_ROWS[-1]})", nf=NF_CHF, italic=True, color="7F7F7F")
        for j, r in enumerate(EXP_ROWS):
            f = fil_switch(sel_fil, [(T_, '"Ausgabe"'), (K_, f"$B{r}")] + crit)
            put(ws, f"{col}{r}", f'=IF($B{r}="",0,{f})', nf=NF_CHF)
        put(ws, f"{col}{R_EXP_TOT}", "=" + fil_switch(sel_fil, [(T_, '"Ausgabe"')] + crit), nf=NF_CHF)
        put(ws, f"{col}{R_EXP_NZ}", f"={col}{R_EXP_TOT}-SUM({col}{EXP_ROWS[0]}:{col}{EXP_ROWS[-1]})", nf=NF_CHF, italic=True, color="7F7F7F")
        put(ws, f"{col}{R_RES}", f"={col}{R_INC_TOT}-{col}{R_EXP_TOT}", nf=NF_CHF)
        put(ws, f"{col}{R_MARGE}", f"=IF({col}{R_INC_TOT}=0,0,{col}{R_RES}/{col}{R_INC_TOT})", nf=NF_PCT)
        put(ws, f"{col}{R_CNT}", "=" + fil_switch(sel_fil, crit, countifs), nf=NF_INT)
    for r in (R_INC_TOT, R_EXP_TOT, R_RES):
        total_style(ws, r, 2, first_col + ncols + 3 if ws.title == "Monatsübersicht" else first_col + ncols - 1)
    return cols


def build_monat(wb, sel_fil_default):
    ws = wb.create_sheet("Monatsübersicht")
    title(ws, "Monatsübersicht – Einnahmen und Ausgaben pro Kategorie (CHF)",
          "Automatisch aus dem Blatt 'Buchungen'. Jahr und Filiale in den gelben Zellen wählen.")
    put(ws, "B4", "Jahr:", bold=True, align="right")
    input_cell(ws, "C4", EXAMPLE_YEAR, nf="0")
    put(ws, "B5", "Filiale:", bold=True, align="right")
    input_cell(ws, "C5", sel_fil_default)
    put(ws, "E4", "Aktive Monate (mit Buchungen):", align="right")
    put(ws, "F4", f'=COUNTIF(C{R_CNT}:N{R_CNT},">0")', nf="0", bold=True)
    put(ws, "E5", "Budgetwerte aus Blatt 'Budget' (Monatswerte der gewählten Filiale)", italic=True, color="7F7F7F")
    put(ws, "B6", "Monat-Nr.", italic=True, color="9C9C9C", align="right")
    for m in range(12):
        col = get_column_letter(3 + m)
        put(ws, f"{col}6", m + 1, italic=True, color="9C9C9C", align="center")
    header_row(ws, 7, 2, ["Kategorie"] + [None] * 12 + ["Total", "Ø pro aktivem Monat", "Anteil %", "Budget / Monat", "Ø Ist − Budget"], height=32)
    for m in range(12):
        col = get_column_letter(3 + m)
        c = ws[f"{col}7"]
        c.value = f"=DATE($C$4,{col}$6,1)"
        c.number_format = NF_MONTH
    build_overview_block(ws, 3, 12, lambda ci: month_crit("$C$4", f"{get_column_letter(3 + ci)}$6"), "$C$5")
    # Zusatzzeilen
    put(ws, f"B{R_CUM}", "Ergebnis kumuliert (Jahresverlauf)", bold=True)
    put(ws, f"B{R_MARGE}", "Marge % (Ergebnis / Einnahmen)")
    put(ws, f"B{R_CNT}", "Anzahl Buchungen")
    put(ws, f"C{R_CUM}", f"=C{R_RES}", nf=NF_CHF, bold=True)
    for m in range(1, 12):
        col, prev = get_column_letter(3 + m), get_column_letter(2 + m)
        put(ws, f"{col}{R_CUM}", f"={prev}{R_CUM}+{col}{R_RES}", nf=NF_CHF, bold=True)
    # Spalten O..S
    all_rows = INC_ROWS + [R_INC_NZ, R_INC_TOT] + EXP_ROWS + [R_EXP_NZ, R_EXP_TOT, R_RES]
    for r in all_rows:
        put(ws, f"O{r}", f"=SUM(C{r}:N{r})", nf=NF_CHF, bold=r in (R_INC_TOT, R_EXP_TOT, R_RES))
        put(ws, f"P{r}", f"=IF($F$4=0,0,O{r}/$F$4)", nf=NF_CHF)
    put(ws, f"O{R_CUM}", f"=N{R_CUM}", nf=NF_CHF, bold=True)
    put(ws, f"O{R_MARGE}", f"=IF(O{R_INC_TOT}=0,0,O{R_RES}/O{R_INC_TOT})", nf=NF_PCT)
    put(ws, f"O{R_CNT}", f"=SUM(C{R_CNT}:N{R_CNT})", nf=NF_INT)
    for r in INC_ROWS + [R_INC_NZ]:
        put(ws, f"Q{r}", f"=IF($O${R_INC_TOT}=0,0,O{r}/$O${R_INC_TOT})", nf=NF_PCT)
    for r in EXP_ROWS + [R_EXP_NZ]:
        put(ws, f"Q{r}", f"=IF($O${R_EXP_TOT}=0,0,O{r}/$O${R_EXP_TOT})", nf=NF_PCT)
    put(ws, f"Q{R_INC_TOT}", "=1", nf=NF_PCT, bold=True)
    put(ws, f"Q{R_EXP_TOT}", "=1", nf=NF_PCT, bold=True)

    def budget_ref(brow):
        return (f'IF($C$5="Alle",Budget!$E${brow},IF($C$5=Listen!$A$3,Budget!$C${brow},'
                f'IF($C$5=Listen!$A$4,Budget!$D${brow},0)))')
    for i, r in enumerate(INC_ROWS):
        put(ws, f"R{r}", "=" + budget_ref(BUD_INC[i]), nf=NF_CHF, color="008000")
    for j, r in enumerate(EXP_ROWS):
        put(ws, f"R{r}", "=" + budget_ref(BUD_EXP[j]), nf=NF_CHF, color="008000")
    put(ws, f"R{R_INC_TOT}", "=" + budget_ref(BUD_INC_TOT), nf=NF_CHF, color="008000", bold=True)
    put(ws, f"R{R_EXP_TOT}", "=" + budget_ref(BUD_EXP_TOT), nf=NF_CHF, color="008000", bold=True)
    put(ws, f"R{R_RES}", "=" + budget_ref(BUD_RES), nf=NF_CHF, color="008000", bold=True)
    for r in INC_ROWS + EXP_ROWS + [R_INC_TOT, R_EXP_TOT, R_RES]:
        put(ws, f"S{r}", f"=P{r}-R{r}", nf=NF_CHF, bold=r in (R_INC_TOT, R_EXP_TOT, R_RES))
    # Abweichung rot: Ausgaben über Budget bzw. Einnahmen/Ergebnis unter Budget
    red = Font(name=FONT, color="C00000", bold=True)
    ws.conditional_formatting.add(f"S{EXP_ROWS[0]}:S{R_EXP_TOT}", FormulaRule(formula=[f"S{EXP_ROWS[0]}>0.005"], font=red))
    ws.conditional_formatting.add(f"S{INC_ROWS[0]}:S{R_INC_TOT}", FormulaRule(formula=[f"S{INC_ROWS[0]}<-0.005"], font=red))
    ws.conditional_formatting.add(f"S{R_RES}", FormulaRule(formula=[f"S{R_RES}<-0.005"], font=red))
    for r in (R_INC_TOT, R_EXP_TOT, R_RES):
        total_style(ws, r, 2, 19)
    put(ws, "B56", "Lesehilfe: 'Nicht zugeordnet' = Buchungen, deren Kategorie nicht (mehr) im Blatt 'Listen' steht – im Journal korrigieren. "
        "'Ø Ist − Budget' positiv bei Ausgaben = über Budget (rot).", italic=True, color="7F7F7F")
    widths(ws, {"A": 2, "B": 46, **{get_column_letter(c): 12.5 for c in range(3, 15)}, "O": 14, "P": 14, "Q": 9, "R": 13, "S": 13})
    ws.freeze_panes = "C8"
    ws.sheet_properties.tabColor = C_HDR
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build_jahr(wb, sel_fil_default):
    ws = wb.create_sheet("Jahresübersicht")
    title(ws, "Jahresübersicht – Mehrjahresvergleich und Jahresdetail (CHF)",
          "Automatisch aus dem Blatt 'Buchungen'. Auswahl in den gelben Zellen.")
    put(ws, "B4", "Filiale:", bold=True, align="right")
    input_cell(ws, "C4", sel_fil_default)
    put(ws, "B5", "Erstes Jahr der Übersicht:", bold=True, align="right")
    input_cell(ws, "C5", EXAMPLE_YEAR, nf="0")
    header_row(ws, 7, 2, ["Kategorie"] + [None] * 6, height=24)
    ws["C7"].value = "=$C$5"
    for i in range(1, 6):
        col, prev = get_column_letter(3 + i), get_column_letter(2 + i)
        ws[f"{col}7"].value = f"={prev}7+1"
    for i in range(6):
        ws[f"{get_column_letter(3 + i)}7"].number_format = "0"
    build_overview_block(ws, 3, 6, lambda ci: year_crit(f"{get_column_letter(3 + ci)}$7"), "$C$4")
    put(ws, f"B{R_MARGE}", "Marge % (Ergebnis / Einnahmen)")
    put(ws, f"B{R_CNT}", "Anzahl Buchungen")
    put(ws, f"B{R_CUM}", "Veränderung Einnahmen zum Vorjahr %")
    put(ws, f"C{R_CUM}", "–", align="center", color="7F7F7F")
    for i in range(1, 6):
        col, prev = get_column_letter(3 + i), get_column_letter(2 + i)
        put(ws, f"{col}{R_CUM}", f"=IF({prev}{R_INC_TOT}=0,0,({col}{R_INC_TOT}-{prev}{R_INC_TOT})/{prev}{R_INC_TOT})", nf=NF_PCT)
    for r in (R_INC_TOT, R_EXP_TOT, R_RES):
        total_style(ws, r, 2, 8)

    # ---- Block 2: Jahr im Detail
    B0 = 58
    put(ws, f"B{B0}", "Jahr im Detail – Filialvergleich und Budget", bold=True, size=13, color=C_HDR)
    put(ws, f"B{B0 + 1}", "Jahr:", bold=True, align="right")
    input_cell(ws, f"C{B0 + 1}", EXAMPLE_YEAR, nf="0")
    put(ws, f"B{B0 + 2}", "Budget-Monate für den Vergleich:", bold=True, align="right")
    input_cell(ws, f"C{B0 + 2}", 12, nf="0")
    ws[f"C{B0 + 2}"].comment = Comment("Anzahl Monate, mit denen das Monatsbudget multipliziert wird. "
                                        "Für ein laufendes Jahr z. B. die Anzahl bereits erfasster Monate eintragen.", "creaDIG")
    hdr = B0 + 4
    header_row(ws, hdr, 2, ["Kategorie", None, None, "Gesamt (alle Filialen)", "Budget", "Abweichung CHF", "Abweichung %"], height=30)
    ws[f"C{hdr}"].value = '=IF(Listen!$A$3="","Filiale 1",Listen!$A$3)'
    ws[f"D{hdr}"].value = '=IF(Listen!$A$4="","Filiale 2",Listen!$A$4)'
    Y = f"$C${B0 + 1}"
    yc = year_crit(Y)
    inc_rows = [hdr + 2 + i for i in range(N_INC)]
    r_inc_nz, r_inc_tot, r_inc_oth = inc_rows[-1] + 1, inc_rows[-1] + 2, inc_rows[-1] + 3
    r_exp_hdr = r_inc_oth + 2
    exp_rows = [r_exp_hdr + 1 + j for j in range(N_EXP)]
    r_exp_nz, r_exp_tot, r_exp_oth = exp_rows[-1] + 1, exp_rows[-1] + 2, exp_rows[-1] + 3
    r_res, r_marge = r_exp_oth + 2, r_exp_oth + 3
    section(ws, hdr + 1, "EINNAHMEN", 8)
    section(ws, r_exp_hdr, "AUSGABEN", 8)
    fil_cols = {"C": "Listen!$A$3", "D": "Listen!$A$4"}

    def cat_block(rows, typ, bud_map, r_nz, r_tot, r_oth, bud_tot, listcol):
        for i, r in enumerate(rows):
            put(ws, f"B{r}", label_link(f"Listen!${listcol}${3 + i}"))
            for col, fref in fil_cols.items():
                put(ws, f"{col}{r}", f'=IF($B{r}="",0,{sumifs((T_, typ), (K_, f"$B{r}"), (F_, fref), *yc)})', nf=NF_CHF)
            put(ws, f"E{r}", f'=IF($B{r}="",0,{sumifs((T_, typ), (K_, f"$B{r}"), *yc)})', nf=NF_CHF)
            put(ws, f"F{r}", f"=Budget!$E${bud_map[i]}*$C${B0 + 2}", nf=NF_CHF, color="008000")
        put(ws, f"B{r_nz}", "Nicht zugeordnet (ohne / unbekannte Kategorie)", italic=True, color="7F7F7F")
        put(ws, f"B{r_tot}", "TOTAL " + ("EINNAHMEN" if typ == '"Einnahme"' else "AUSGABEN"), bold=True)
        put(ws, f"B{r_oth}", "davon andere / keine Filiale (in 'Gesamt' enthalten)", italic=True, color="7F7F7F")
        for col, fref in fil_cols.items():
            put(ws, f"{col}{r_tot}", "=" + sumifs((T_, typ), (F_, fref), *yc), nf=NF_CHF)
            put(ws, f"{col}{r_nz}", f"={col}{r_tot}-SUM({col}{rows[0]}:{col}{rows[-1]})", nf=NF_CHF, italic=True, color="7F7F7F")
        put(ws, f"E{r_tot}", "=" + sumifs((T_, typ), *yc), nf=NF_CHF)
        put(ws, f"E{r_nz}", f"=E{r_tot}-SUM(E{rows[0]}:E{rows[-1]})", nf=NF_CHF, italic=True, color="7F7F7F")
        put(ws, f"E{r_oth}", f"=E{r_tot}-C{r_tot}-D{r_tot}", nf=NF_CHF, italic=True, color="7F7F7F")
        put(ws, f"F{r_tot}", f"=Budget!$E${bud_tot}*$C${B0 + 2}", nf=NF_CHF, color="008000", bold=True)
        for r in rows + [r_tot]:
            put(ws, f"G{r}", f"=E{r}-F{r}", nf=NF_CHF, bold=(r == r_tot))
            put(ws, f"H{r}", f"=IF(F{r}=0,0,G{r}/F{r})", nf=NF_PCT, bold=(r == r_tot))
        total_style(ws, r_tot, 2, 8)

    cat_block(inc_rows, '"Einnahme"', BUD_INC, r_inc_nz, r_inc_tot, r_inc_oth, BUD_INC_TOT, "B")
    cat_block(exp_rows, '"Ausgabe"', BUD_EXP, r_exp_nz, r_exp_tot, r_exp_oth, BUD_EXP_TOT, "C")
    put(ws, f"B{r_res}", "ERGEBNIS (Einnahmen − Ausgaben)", bold=True)
    put(ws, f"B{r_marge}", "Marge % (Ergebnis / Einnahmen)")
    for col in "CDE":
        put(ws, f"{col}{r_res}", f"={col}{r_inc_tot}-{col}{r_exp_tot}", nf=NF_CHF, bold=True)
        put(ws, f"{col}{r_marge}", f"=IF({col}{r_inc_tot}=0,0,{col}{r_res}/{col}{r_inc_tot})", nf=NF_PCT)
    put(ws, f"F{r_res}", f"=Budget!$E${BUD_RES}*$C${B0 + 2}", nf=NF_CHF, color="008000", bold=True)
    put(ws, f"G{r_res}", f"=E{r_res}-F{r_res}", nf=NF_CHF, bold=True)
    put(ws, f"H{r_res}", f"=IF(F{r_res}=0,0,G{r_res}/F{r_res})", nf=NF_PCT, bold=True)
    put(ws, f"F{r_marge}", f"=Budget!$F${BUD_MARGE}", nf=NF_PCT, color="008000")
    total_style(ws, r_res, 2, 8)
    red = Font(name=FONT, color="C00000", bold=True)
    ws.conditional_formatting.add(f"G{exp_rows[0]}:H{r_exp_tot}", FormulaRule(formula=[f"$G{exp_rows[0]}>0.005"], font=red))
    ws.conditional_formatting.add(f"G{inc_rows[0]}:H{r_inc_tot}", FormulaRule(formula=[f"$G{inc_rows[0]}<-0.005"], font=red))
    ws.conditional_formatting.add(f"G{r_res}:H{r_res}", FormulaRule(formula=[f"$G{r_res}<-0.005"], font=red))
    widths(ws, {"A": 2, "B": 46, "C": 14, "D": 14, "E": 16, "F": 14, "G": 14, "H": 12})
    ws.freeze_panes = "C8"
    ws.sheet_properties.tabColor = C_HDR
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws._detail_rows = {"inc": inc_rows, "inc_tot": r_inc_tot, "exp": exp_rows, "exp_tot": r_exp_tot, "res": r_res, "hdr": hdr}
    return ws


def build_mitarbeiter(wb):
    ws = wb.create_sheet("Mitarbeiter")
    title(ws, "Mitarbeiter – Personalkosten pro Filiale (CHF)",
          "Planungsblatt: Bruttolöhne und geschätzte Arbeitgeberbeiträge. Die effektiv bezahlten Löhne werden zusätzlich im Journal erfasst.")
    put(ws, "B4", "Arbeitgeberbeiträge in % des Bruttolohns (Schätzwerte – bitte mit Treuhänder / Lohnabrechnung abgleichen)", bold=True)
    for i, (lab, val) in enumerate(AG_SAETZE):
        put(ws, f"B{5 + i}", lab)
        input_cell(ws, f"C{5 + i}", val, nf="0.0%")
    put(ws, "B10", "Total Arbeitgeberbeiträge", bold=True)
    put(ws, "C10", "=SUM(C5:C9)", nf="0.0%", bold=True, border=BORDER)
    ws["C5"].comment = Comment("AHV/IV/EO: Arbeitgeberanteil 5.3 % (Gesamtsatz 10.6 %). ALV: 1.1 % bis CHF 148'200 Jahreslohn. "
                               "BVG, UVG/KTG und FAK sind vertrags- bzw. kantonsabhängig – Schätzwerte.", "creaDIG")
    put(ws, "E4", "Legende: gelb = Eingabe · schwarz = Formel. Pensum in % (100 % = Vollzeit). "
        "Bruttolohn = effektiver Monatslohn für das Pensum (bei Stundenlohn: Durchschnitt). Beispielzeilen ersetzen.", italic=True, color="7F7F7F", wrap=True)
    ws.merge_cells("E4:K6")
    header_row(ws, 13, 2, ["Name", "Filiale", "Funktion", "Eintritt", "Pensum %", "Bruttolohn / Monat CHF", "13. Monatslohn",
                           "AG-Beiträge / Monat CHF", "Personalkosten / Monat CHF", "Personalkosten / Jahr CHF"], height=34)
    for k in range(MA_N):
        r = MA_FIRST + k
        ex = MITARBEITER_BEISPIELE[k] if k < len(MITARBEITER_BEISPIELE) else (None,) * 7
        input_cell(ws, f"B{r}", ex[0], align="left")
        input_cell(ws, f"C{r}", ex[1])
        input_cell(ws, f"D{r}", ex[2], align="left")
        input_cell(ws, f"E{r}", ex[3], nf=NF_DATE)
        input_cell(ws, f"F{r}", ex[4], nf="0%")
        input_cell(ws, f"G{r}", ex[5], nf=NF_CHF)
        input_cell(ws, f"H{r}", ex[6])
        put(ws, f"I{r}", f'=IF($G{r}="",0,$G{r}*IF($H{r}="Ja",13,12)/12*$C$10)', nf=NF_CHF, border=BORDER)
        put(ws, f"J{r}", f'=IF($G{r}="",0,$G{r}*IF($H{r}="Ja",13,12)/12+I{r})', nf=NF_CHF, border=BORDER)
        put(ws, f"K{r}", f"=J{r}*12", nf=NF_CHF, border=BORDER)
    dv_fil = DataValidation(type="list", formula1=f"OFFSET(Listen!$A$3,0,0,MAX(1,COUNTA(Listen!$A$3:$A${2 + N_FIL})),1)", allow_blank=True)
    dv_fil.add(f"C{MA_FIRST}:C{MA_LAST}")
    dv_13 = DataValidation(type="list", formula1='"Ja,Nein"', allow_blank=True)
    dv_13.add(f"H{MA_FIRST}:H{MA_LAST}")
    dv_pens = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True, showErrorMessage=True,
                             errorTitle="Pensum", error="Pensum als Prozentwert zwischen 0 % und 100 % eingeben.")
    dv_pens.add(f"F{MA_FIRST}:F{MA_LAST}")
    for dv in (dv_fil, dv_13, dv_pens):
        ws.add_data_validation(dv)
    put(ws, "B36", "Zusammenfassung pro Filiale", bold=True, size=12, color=C_HDR)
    header_row(ws, 37, 2, ["Filiale", "Anzahl Mitarbeitende", "Pensum total", "Bruttolöhne Ø / Monat (inkl. 13.)",
                           "AG-Beiträge / Monat", "Personalkosten / Monat", "Personalkosten / Jahr"], height=34)
    rows_ = f"${MA_FIRST}:$" 
    C, F, G, H, I, J = (f"${c}${MA_FIRST}:${c}${MA_LAST}" for c in "CFGHIJ")
    for name, r in MA_SUM.items():
        if name == "Gesamt":
            put(ws, f"B{r}", "Gesamt", bold=True)
            for col in "CDEFGH":
                put(ws, f"{col}{r}", f"=SUM({col}{MA_SUM['Zürich']}:{col}{MA_SUM['Luzern']})", nf=NF_CHF if col not in "CD" else ("0" if col == "C" else "0%"), bold=True)
            total_style(ws, r, 2, 8)
            continue
        put(ws, f"B{r}", '=IF(Listen!$A$3="","Filiale 1",Listen!$A$3)' if name == "Zürich" else '=IF(Listen!$A$4="","Filiale 2",Listen!$A$4)', bold=True)
        put(ws, f"C{r}", f'=COUNTIFS({C},$B{r},{G},">0")', nf="0")
        put(ws, f"D{r}", f"=SUMIFS({F},{C},$B{r})", nf="0%")
        put(ws, f"E{r}", f'=SUMPRODUCT(({C}=$B{r})*({G})*(12+({H}="Ja"))/12)', nf=NF_CHF)
        put(ws, f"F{r}", f"=SUMIFS({I},{C},$B{r})", nf=NF_CHF)
        put(ws, f"G{r}", f"=SUMIFS({J},{C},$B{r})", nf=NF_CHF)
        put(ws, f"H{r}", f"=G{r}*12", nf=NF_CHF)
    put(ws, "B42", "Diese Werte fliessen als Planwerte in das Blatt 'Budget' (Zeilen 'Löhne & Gehälter' und 'Sozialversicherungen') "
        "und dienen dem Vergleich mit den tatsächlich im Journal erfassten Lohnzahlungen.", italic=True, color="7F7F7F")
    widths(ws, {"A": 2, "B": 30, "C": 14, "D": 26, "E": 14, "F": 12, "G": 16, "H": 14, "I": 16, "J": 18, "K": 18})
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "B14"
    ws.sheet_properties.tabColor = "5B9BD5"
    return ws


def build_budget(wb):
    ws = wb.create_sheet("Budget")
    title(ws, "Budget – monatliche Planwerte pro Kategorie und Filiale (CHF)",
          "Gelb = Eingabe (Planwert pro Monat). Vorbelegte Zahlen sind Beispiel-Annahmen – bitte durch eigene Planwerte ersetzen. "
          "Grün = automatisch aus 'Mitarbeiter' (überschreibbar).")
    header_row(ws, 4, 2, ["Kategorie", None, None, "Gesamt / Monat", "Gesamt / Jahr"], height=28)
    ws["C4"].value = '=IF(Listen!$A$3="","Filiale 1",Listen!$A$3)'
    ws["D4"].value = '=IF(Listen!$A$4="","Filiale 2",Listen!$A$4)'
    section(ws, 5, "EINNAHMEN", 6)
    section(ws, 18, "AUSGABEN", 6)
    # Beispiel-Planwerte (Monat)
    inc_plan = {0: (10800, 7200), 1: (32400, 21600), 2: (9200, 6100), 3: (1600, 1000), 4: (500, 300)}
    exp_plan = {0: (15100, 10100), 1: (2200, 1450), 2: (1600, 1050), 5: (5800, 3400), 6: (720, 480), 7: (240, 190),
                8: (700, 450), 9: (2800, 1850), 10: (450, 300), 11: (380, 260), 12: (250, 150), 13: (300, 150),
                14: (285, 215), 15: (1000, 650), 16: (350, 250), 17: (320, 0), 18: (45, 35), 19: (150, 100)}
    for i in range(N_INC):
        r = BUD_INC[i]
        put(ws, f"B{r}", label_link(f"Listen!$B${3 + i}"))
        zh, lu = inc_plan.get(i, (None, None))
        input_cell(ws, f"C{r}", zh, nf=NF_CHF, align="right")
        input_cell(ws, f"D{r}", lu, nf=NF_CHF, align="right")
    for j in range(N_EXP):
        r = BUD_EXP[j]
        put(ws, f"B{r}", label_link(f"Listen!$C${3 + j}"))
        if j == 3:    # Löhne & Gehälter aus Mitarbeiter
            put(ws, f"C{r}", f"=Mitarbeiter!$E${MA_SUM['Zürich']}", nf=NF_CHF, color="008000", fill_=C_IN, border=BORDER)
            put(ws, f"D{r}", f"=Mitarbeiter!$E${MA_SUM['Luzern']}", nf=NF_CHF, color="008000", fill_=C_IN, border=BORDER)
        elif j == 4:  # Sozialversicherungen aus Mitarbeiter
            put(ws, f"C{r}", f"=Mitarbeiter!$F${MA_SUM['Zürich']}", nf=NF_CHF, color="008000", fill_=C_IN, border=BORDER)
            put(ws, f"D{r}", f"=Mitarbeiter!$F${MA_SUM['Luzern']}", nf=NF_CHF, color="008000", fill_=C_IN, border=BORDER)
        else:
            zh, lu = exp_plan.get(j, (None, None))
            input_cell(ws, f"C{r}", zh, nf=NF_CHF, align="right")
            input_cell(ws, f"D{r}", lu, nf=NF_CHF, align="right")
    for r in list(BUD_INC.values()) + list(BUD_EXP.values()):
        put(ws, f"E{r}", f"=SUM(C{r}:D{r})", nf=NF_CHF)
        put(ws, f"F{r}", f"=E{r}*12", nf=NF_CHF)
    put(ws, f"B{BUD_INC_TOT}", "TOTAL EINNAHMEN (Plan)", bold=True)
    put(ws, f"B{BUD_EXP_TOT}", "TOTAL AUSGABEN (Plan)", bold=True)
    put(ws, f"B{BUD_RES}", "ERGEBNIS (Plan)", bold=True)
    put(ws, f"B{BUD_MARGE}", "Marge % (Plan)")
    for col in "CDEF":
        put(ws, f"{col}{BUD_INC_TOT}", f"=SUM({col}{BUD_INC[0]}:{col}{BUD_INC[N_INC - 1]})", nf=NF_CHF, bold=True)
        put(ws, f"{col}{BUD_EXP_TOT}", f"=SUM({col}{BUD_EXP[0]}:{col}{BUD_EXP[N_EXP - 1]})", nf=NF_CHF, bold=True)
        put(ws, f"{col}{BUD_RES}", f"={col}{BUD_INC_TOT}-{col}{BUD_EXP_TOT}", nf=NF_CHF, bold=True)
        put(ws, f"{col}{BUD_MARGE}", f"=IF({col}{BUD_INC_TOT}=0,0,{col}{BUD_RES}/{col}{BUD_INC_TOT})", nf=NF_PCT)
    for r in (BUD_INC_TOT, BUD_EXP_TOT, BUD_RES):
        total_style(ws, r, 2, 6)
    put(ws, "B49", "Quelle der Beispielwerte: grobe Annahmen für einen veganen Takeaway (Wareneinsatz ca. 30 %, Personal ca. 30 %, Miete ca. 12 % des Umsatzes). "
        "Keine echten Vegitat-Zahlen.", italic=True, color="7F7F7F")
    widths(ws, {"A": 2, "B": 46, "C": 15, "D": 15, "E": 16, "F": 16})
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "C5"
    ws.sheet_properties.tabColor = "5B9BD5"
    return ws


def build_dashboard(wb):
    ws = wb.create_sheet("Dashboard", 1)
    put(ws, "B1", "VEGITAT – Finanz-Dashboard", bold=True, size=18, color=C_HDR)
    put(ws, "B2", "Zürich · Luzern · alle Beträge in CHF · rechnet sich automatisch aus dem Blatt 'Buchungen'", italic=True, color="595959")
    ws.row_dimensions[1].height = 28
    put(ws, "B4", "Auswertungsjahr:", bold=True, align="right")
    input_cell(ws, "C4", EXAMPLE_YEAR, nf="0")
    put(ws, "E4", "Letzte Buchung:", align="right")
    put(ws, "F4", f'=IF(COUNT({D_})=0,"–",MAX({D_}))', nf=NF_DATE, bold=True, align="center")
    put(ws, "H4", "Buchungen gesamt (alle Jahre):", align="right")
    put(ws, "I4", f"=COUNT({B_})", nf="#,##0", bold=True, align="center")
    Y = "$C$4"
    yc = year_crit(Y)
    # ---------- Filialtabelle (Zeilen 10-15)
    put(ws, "B10", '="Filialvergleich "&$C$4', bold=True, size=12, color=C_HDR)
    header_row(ws, 11, 2, ["Filiale", "Einnahmen", "Ausgaben", "Ergebnis", "Marge %", "Anteil Einnahmen %", "Buchungen"], height=28)
    put(ws, "B12", '=IF(Listen!$A$3="","Filiale 1",Listen!$A$3)', bold=True)
    put(ws, "B13", '=IF(Listen!$A$4="","Filiale 2",Listen!$A$4)', bold=True)
    put(ws, "B14", "Andere / ohne Filiale", italic=True, color="7F7F7F")
    put(ws, "B15", "Gesamt", bold=True)
    for r in (12, 13):
        put(ws, f"C{r}", "=" + sumifs((T_, '"Einnahme"'), (F_, f"$B{r}"), *yc), nf=NF_CHF)
        put(ws, f"D{r}", "=" + sumifs((T_, '"Ausgabe"'), (F_, f"$B{r}"), *yc), nf=NF_CHF)
        put(ws, f"H{r}", "=" + countifs((F_, f"$B{r}"), *yc), nf=NF_INT)
    put(ws, "C15", "=" + sumifs((T_, '"Einnahme"'), *yc), nf=NF_CHF, bold=True)
    put(ws, "D15", "=" + sumifs((T_, '"Ausgabe"'), *yc), nf=NF_CHF, bold=True)
    put(ws, "H15", "=" + countifs(*yc), nf=NF_INT, bold=True)
    for col in "CDH":
        put(ws, f"{col}14", f"={col}15-{col}12-{col}13", nf=NF_CHF if col != "H" else NF_INT, italic=True, color="7F7F7F")
    for r in (12, 13, 14, 15):
        put(ws, f"E{r}", f"=C{r}-D{r}", nf=NF_CHF, bold=(r == 15))
        put(ws, f"F{r}", f"=IF(C{r}=0,0,E{r}/C{r})", nf=NF_PCT)
        put(ws, f"G{r}", f"=IF($C$15=0,0,C{r}/$C$15)", nf=NF_PCT)
    total_style(ws, 15, 2, 8)

    # ---------- Monatstabelle (Zeilen 72-86)
    T0 = 72
    put(ws, f"B{T0}", '="Monatswerte "&$C$4&" (Datenbasis der Diagramme)"', bold=True, size=12, color=C_HDR)
    header_row(ws, T0 + 1, 1, ["Nr.", "Monat", "Einnahmen", "Ausgaben", "Ergebnis", "Ergebnis kumuliert", None, None, None, None, "Buchungen"], height=30)
    ws[f"G{T0 + 1}"].value = '=IF(Listen!$A$3="","Filiale 1",Listen!$A$3)&" – Einnahmen"'
    ws[f"H{T0 + 1}"].value = '=IF(Listen!$A$4="","Filiale 2",Listen!$A$4)&" – Einnahmen"'
    ws[f"I{T0 + 1}"].value = '=IF(Listen!$A$3="","Filiale 1",Listen!$A$3)&" – Ergebnis"'
    ws[f"J{T0 + 1}"].value = '=IF(Listen!$A$4="","Filiale 2",Listen!$A$4)&" – Ergebnis"'
    m_first, m_last = T0 + 2, T0 + 13
    for m in range(12):
        r = m_first + m
        put(ws, f"A{r}", m + 1, color="9C9C9C", align="center")
        put(ws, f"B{r}", MONATE[m], bold=True)
        mc = month_crit(Y, f"$A{r}")
        put(ws, f"C{r}", "=" + sumifs((T_, '"Einnahme"'), *mc), nf=NF_CHF)
        put(ws, f"D{r}", "=" + sumifs((T_, '"Ausgabe"'), *mc), nf=NF_CHF)
        put(ws, f"E{r}", f"=C{r}-D{r}", nf=NF_CHF)
        put(ws, f"F{r}", f"=E{r}" if m == 0 else f"=F{r - 1}+E{r}", nf=NF_CHF)
        put(ws, f"G{r}", "=" + sumifs((T_, '"Einnahme"'), (F_, "Listen!$A$3"), *mc), nf=NF_CHF)
        put(ws, f"H{r}", "=" + sumifs((T_, '"Einnahme"'), (F_, "Listen!$A$4"), *mc), nf=NF_CHF)
        put(ws, f"I{r}", f"=G{r}-" + sumifs((T_, '"Ausgabe"'), (F_, "Listen!$A$3"), *mc), nf=NF_CHF)
        put(ws, f"J{r}", f"=H{r}-" + sumifs((T_, '"Ausgabe"'), (F_, "Listen!$A$4"), *mc), nf=NF_CHF)
        put(ws, f"K{r}", "=" + countifs(*mc), nf=NF_INT)
    r_tot = m_last + 1
    put(ws, f"B{r_tot}", "Total", bold=True)
    for col in "CDEGHIJK":
        put(ws, f"{col}{r_tot}", f"=SUM({col}{m_first}:{col}{m_last})", nf=NF_CHF if col != "K" else NF_INT, bold=True)
    put(ws, f"F{r_tot}", f"=F{m_last}", nf=NF_CHF, bold=True)
    total_style(ws, r_tot, 2, 11)

    # ---------- Ausgaben-Kategorien: Liste (L88..) und Top 10 (L74..)
    C0 = T0 + 14   # 86
    put(ws, f"M{C0}", '="Alle Ausgaben-Kategorien "&$C$4', bold=True, size=12, color=C_HDR)
    header_row(ws, C0 + 1, 12, ["Nr.", "Kategorie", "Betrag", "Sortierhilfe"], height=30)
    cat_first, cat_last = C0 + 2, C0 + 1 + N_EXP
    for j in range(N_EXP):
        r = cat_first + j
        put(ws, f"L{r}", j + 1, color="9C9C9C", align="center")
        put(ws, f"M{r}", label_link(f"Listen!$C${3 + j}"))
        put(ws, f"N{r}", f'=IF($M{r}="",0,{sumifs((T_, chr(34) + "Ausgabe" + chr(34)), (K_, f"$M{r}"), *yc)})', nf=NF_CHF)
        put(ws, f"O{r}", f'=IF($M{r}="",-1,N{r})+ROW()/1000000000', nf="0.00", color="9C9C9C")
    put(ws, f"M{T0}", '="Top 10 Ausgaben-Kategorien "&$C$4', bold=True, size=12, color=C_HDR)
    header_row(ws, T0 + 1, 12, ["Rang", "Kategorie", "Betrag", "Anteil %"], height=30)
    top_first = T0 + 2
    for k in range(10):
        r = top_first + k
        big = f"LARGE($O${cat_first}:$O${cat_last},{k + 1})"
        put(ws, f"L{r}", k + 1, align="center")
        put(ws, f"M{r}", f'=IF({big}<0.01,"",INDEX($M${cat_first}:$M${cat_last},MATCH({big},$O${cat_first}:$O${cat_last},0)))')
        put(ws, f"N{r}", f'=IF({big}<0.01,0,INDEX($N${cat_first}:$N${cat_last},MATCH({big},$O${cat_first}:$O${cat_last},0)))', nf=NF_CHF)
        put(ws, f"O{r}", f"=IF($D$15=0,0,N{r}/$D$15)", nf=NF_PCT)

    # ---------- Zahlungsarten (Zeilen 89-101)
    P0 = T0 + 17   # 89
    put(ws, f"B{P0}", '="Einnahmen nach Zahlungsart "&$C$4', bold=True, size=12, color=C_HDR)
    header_row(ws, P0 + 1, 2, ["Zahlungsart", "Betrag", "Anteil %"], height=30)
    pay_first, pay_last = P0 + 2, P0 + 1 + N_PAY
    for p in range(N_PAY):
        r = pay_first + p
        put(ws, f"B{r}", label_link(f"Listen!$D${3 + p}"))
        put(ws, f"C{r}", f'=IF($B{r}="",0,{sumifs((T_, chr(34) + "Einnahme" + chr(34)), (Z_, f"$B{r}"), *yc)})', nf=NF_CHF)
        put(ws, f"D{r}", f"=IF($C$15=0,0,C{r}/$C$15)", nf=NF_PCT)
    r_oth = pay_last + 1
    put(ws, f"B{r_oth}", "Ohne / unbekannte Zahlungsart", italic=True, color="7F7F7F")
    put(ws, f"C{r_oth}", f"=$C$15-SUM(C{pay_first}:C{pay_last})", nf=NF_CHF, italic=True, color="7F7F7F")
    put(ws, f"D{r_oth}", f"=IF($C$15=0,0,C{r_oth}/$C$15)", nf=NF_PCT, italic=True, color="7F7F7F")

    # ---------- Kacheln (Zeilen 6-8)
    active = f'COUNTIF($K${m_first}:$K${m_last},">0")'
    tiles = [
        ("B", "Einnahmen", "=$C$15", NF_CHF0, "Ø pro aktivem Monat", f"=IF({active}=0,0,$C$15/{active})", NF_CHF0),
        ("E", "Ausgaben", "=$D$15", NF_CHF0, "Ø pro aktivem Monat", f"=IF({active}=0,0,$D$15/{active})", NF_CHF0),
        ("H", "Ergebnis (Gewinn / Verlust)", "=$E$15", NF_CHF0, "Ø pro aktivem Monat", f"=IF({active}=0,0,$E$15/{active})", NF_CHF0),
        ("K", "Marge (Ergebnis / Einnahmen)", "=$F$15", NF_PCT, "Plan-Marge laut Budget", f"=Budget!$F${BUD_MARGE}", NF_PCT),
        ("N", "Bester Monat (Ergebnis)",
         f'=IF(MAX($E${m_first}:$E${m_last})<=0,"–",INDEX($B${m_first}:$B${m_last},MATCH(MAX($E${m_first}:$E${m_last}),$E${m_first}:$E${m_last},0)))',
         "@", "Ergebnis in diesem Monat", f"=MAX($E${m_first}:$E${m_last})", NF_CHF0),
        ("Q", "Unvollständige Buchungen (alle Jahre)",
         f'=SUMPRODUCT(({B_}<>"")*((({D_}="")+({F_}="")+({T_}="")+({K_}=""))>0))',
         "#,##0", "davon mit unbekannter Kategorie",
         f'=SUMPRODUCT(({B_}<>"")*({K_}<>"")*(COUNTIF(Listen!$B$3:$C${2 + N_EXP},{K_})=0))', "#,##0"),
    ]
    for col, lab, f, nf, sublab, subf, subnf in tiles:
        nxt = get_column_letter(ord(col) - 64 + 1)
        for r in (6, 7, 8):
            for c in (col, nxt):
                ws[f"{c}{r}"].fill = fill(C_TILE)
        ws.merge_cells(f"{col}6:{nxt}6")
        ws.merge_cells(f"{col}7:{nxt}7")
        put(ws, f"{col}6", lab, bold=True, color="595959", align="center", fill_=C_TILE)
        put(ws, f"{col}7", f, bold=True, size=16, color=C_HDR, nf=nf, align="center", fill_=C_TILE)
        put(ws, f"{col}8", sublab, size=8, color="7F7F7F", align="left", fill_=C_TILE)
        put(ws, f"{nxt}8", subf, size=8, bold=True, color="595959", nf=subnf, align="right", fill_=C_TILE)
    ws.row_dimensions[7].height = 30
    ws.conditional_formatting.add("Q7:R7", FormulaRule(formula=["$Q$7>0"], font=Font(name=FONT, bold=True, size=16, color="C00000")))
    put(ws, "B9", "Hinweis: 'Unvollständig' = Zeilen mit Betrag, aber ohne Datum, Filiale, Typ oder Kategorie (im Journal orange markiert). "
        "'Unbekannte Kategorie' = Kategorie-Name steht nicht im Blatt 'Listen'.", size=8, italic=True, color="7F7F7F")

    # ---------- Diagramme
    def style_bar(ch, colors, horizontal=False):
        ch.type = "bar" if horizontal else "col"
        ch.grouping = "clustered"
        ch.gapWidth = 60
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.numFmt = "#,##0"
        ch.y_axis.title = "CHF"
        ch.legend.position = "b"
        ch.width, ch.height = 20, 8.5
        for s, colr in zip(ch.series, colors):
            s.graphicalProperties.solidFill = colr
            s.graphicalProperties.line.solidFill = colr

    def style_line(ch, colors):
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.numFmt = "#,##0"
        ch.y_axis.title = "CHF"
        ch.legend.position = "b"
        ch.width, ch.height = 20, 8.5
        for s, colr in zip(ch.series, colors):
            s.graphicalProperties.line.solidFill = colr
            s.graphicalProperties.line.width = 25000
            s.marker.symbol = "circle"
            s.marker.size = 6
            s.marker.graphicalProperties.solidFill = colr
            s.marker.graphicalProperties.line.solidFill = colr
            s.smooth = False

    cats = Reference(ws, min_col=2, min_row=m_first, max_row=m_last)
    c1 = BarChart()
    c1.title = "Einnahmen vs. Ausgaben pro Monat"
    c1.add_data(Reference(ws, min_col=3, max_col=4, min_row=T0 + 1, max_row=m_last), titles_from_data=True)
    c1.set_categories(cats)
    style_bar(c1, [COL_BLUE, COL_ORANGE])
    ws.add_chart(c1, "B17")

    c2 = LineChart()
    c2.title = "Ergebnis pro Monat und kumuliert"
    c2.add_data(Reference(ws, min_col=5, max_col=6, min_row=T0 + 1, max_row=m_last), titles_from_data=True)
    c2.set_categories(cats)
    style_line(c2, [COL_AQUA, COL_VIOLET])
    ws.add_chart(c2, "L17")

    c3 = BarChart()
    c3.title = "Top 10 Ausgaben-Kategorien"
    c3.add_data(Reference(ws, min_col=14, min_row=T0 + 1, max_row=top_first + 9), titles_from_data=True)
    c3.set_categories(Reference(ws, min_col=13, min_row=top_first, max_row=top_first + 9))
    style_bar(c3, [COL_ORANGE], horizontal=True)
    c3.x_axis.scaling.orientation = "maxMin"
    c3.legend = None
    ws.add_chart(c3, "B35")

    c4 = BarChart()
    c4.title = "Einnahmen je Filiale pro Monat"
    c4.add_data(Reference(ws, min_col=7, max_col=8, min_row=T0 + 1, max_row=m_last), titles_from_data=True)
    c4.set_categories(cats)
    style_bar(c4, [COL_AQUA, COL_VIOLET])
    ws.add_chart(c4, "L35")

    c5 = LineChart()
    c5.title = "Ergebnis je Filiale pro Monat"
    c5.add_data(Reference(ws, min_col=9, max_col=10, min_row=T0 + 1, max_row=m_last), titles_from_data=True)
    c5.set_categories(cats)
    style_line(c5, [COL_AQUA, COL_VIOLET])
    ws.add_chart(c5, "B53")

    c6 = BarChart()
    c6.title = "Einnahmen nach Zahlungsart"
    c6.add_data(Reference(ws, min_col=3, min_row=P0 + 1, max_row=pay_last), titles_from_data=True)
    c6.set_categories(Reference(ws, min_col=2, min_row=pay_first, max_row=pay_last))
    style_bar(c6, [COL_BLUE], horizontal=True)
    c6.x_axis.scaling.orientation = "maxMin"
    c6.legend = None
    ws.add_chart(c6, "L53")

    widths(ws, {"A": 4, **{get_column_letter(c): 12.5 for c in range(2, 21)}})
    ws.column_dimensions["M"].width = 40
    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = "C00000"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws._layout = {"m_first": m_first, "m_last": m_last, "cat_first": cat_first, "top_first": top_first, "pay_first": pay_first, "pay_oth": r_oth}
    return ws


def build_anleitung(wb):
    ws = wb.create_sheet("Anleitung", 0)
    put(ws, "B1", "Vegitat – Finanzübersicht Zürich & Luzern (CHF)", bold=True, size=18, color=C_HDR)
    put(ws, "B2", "Einnahmen und Ausgaben erfassen – Monats- und Jahresauswertung mit Diagrammen, automatisch.", italic=True, color="595959")
    ws.row_dimensions[1].height = 28
    lines = [
        ("H", "So ist die Datei aufgebaut"),
        ("T", "1. Buchungen – das einzige Blatt, das täglich gepflegt wird: jede Einnahme und jede Ausgabe als eigene Zeile (Datum, Filiale, Typ, Kategorie, Beschreibung, Betrag, Zahlungsart)."),
        ("T", "2. Dashboard – Kennzahlen, Filialvergleich und 6 Diagramme für das gewählte Jahr. Rechnet sich automatisch."),
        ("T", "3. Monatsübersicht – alle Kategorien × 12 Monate, wahlweise für alle Filialen oder eine einzelne, inklusive Budgetvergleich."),
        ("T", "4. Jahresübersicht – Vergleich über 6 Jahre sowie ein Jahr im Detail: Zürich | Luzern | Gesamt | Budget | Abweichung."),
        ("T", "5. Mitarbeiter – Personalliste mit Lohnkosten inklusive Arbeitgeberbeiträgen pro Filiale (Planung). Fliesst ins Budget."),
        ("T", "6. Budget – monatliche Planwerte je Kategorie und Filiale."),
        ("T", "7. Listen – Filialen, Kategorien, Zahlungsarten. Hier anpassen; alle Dropdowns und Auswertungen folgen automatisch."),
        ("", ""),
        ("H", "Farblegende"),
        ("L1", "Gelbe Zelle = Eingabe oder Auswahl (z. B. Jahr, Filiale, Planwerte)."),
        ("L2", "Blaue Schrift = eingegebener Wert · schwarze Schrift = Formel (nicht überschreiben) · grüne Schrift = Verknüpfung aus einem anderen Blatt."),
        ("L3", "Orange Zeile im Journal = unvollständige Buchung (Betrag vorhanden, aber Datum, Filiale, Typ oder Kategorie fehlt)."),
        ("", ""),
        ("H", "Regeln für die Erfassung im Blatt 'Buchungen'"),
        ("T", "• Betrag immer positiv und in CHF inklusive MWST (Bruttobetrag). Ob Einnahme oder Ausgabe, entscheidet die Spalte 'Typ'."),
        ("T", "• Tagesumsatz: pro Tag und Filiale eine Zeile je Zahlungsart, z. B. 'Tagesumsatz Kasse Bar' und 'Tagesumsatz Karte/Twint'. So bleibt das Bargeld nachvollziehbar."),
        ("T", "• Lieferdienste (Uber Eats, Just Eat, Smood): Bruttobestellwert als Einnahme 'Lieferdienste' und die Kommission separat als Ausgabe 'Lieferplattform-Gebühren'. So sieht man, was die Plattformen wirklich kosten."),
        ("T", "• Löhne: monatlich die ausbezahlten Bruttolöhne als 'Löhne & Gehälter' und die Beiträge an AHV/ALV/BVG/UVG als 'Sozialversicherungen' erfassen."),
        ("T", "• Rückerstattungen: als Einnahme 'Sonstige Einnahmen' bzw. Ausgabe 'Sonstige Ausgaben' mit Beschreibung 'Rückerstattung …'."),
        ("T", "• Privatbezüge, Einlagen und Kreditrückzahlungen sind keine Betriebsausgaben – entweder nicht erfassen oder eine eigene Kategorie dafür anlegen."),
        ("T", "• Neue Buchung: in die nächste leere Zeile direkt unter der Tabelle schreiben. Formatierung und Dropdowns erweitern sich automatisch (bis Zeile 5'000)."),
        ("T", "• Beispielzeilen löschen: im Journal Filter auf Bemerkung = 'Beispiel' setzen, alle gefilterten Zeilen markieren, Rechtsklick → 'Tabellenzeilen löschen'. Danach im Dashboard das Jahr auf 2026 stellen."),
        ("", ""),
        ("H", "Auswerten"),
        ("T", "• Dashboard, Monats- und Jahresübersicht: gelbe Zellen 'Jahr' und 'Filiale' ändern – alles rechnet sofort neu, auch die Diagramme."),
        ("T", "• 'Nicht zugeordnet' = Buchungen, deren Kategorie nicht (mehr) im Blatt 'Listen' steht. Solche Zeilen im Journal korrigieren."),
        ("T", "• Die Kachel 'Unvollständige Buchungen' im Dashboard zeigt, ob im Journal etwas fehlt. Ziel: 0."),
        ("T", "• Budgetvergleich: 'Ø Ist − Budget' bzw. 'Abweichung' in Rot bedeutet: Ausgaben über Plan oder Einnahmen unter Plan."),
        ("", ""),
        ("H", "Anpassen"),
        ("T", "• Filiale, Kategorie oder Zahlungsart hinzufügen: im Blatt 'Listen' in die nächste freie Zeile schreiben (max. 10 Filialen, 10 Einnahme- und 25 Ausgabe-Kategorien)."),
        ("T", "• Kategorie umbenennen: zuerst im Journal per Suchen/Ersetzen alle alten Einträge umbenennen, dann in 'Listen' – sonst laufen alte Buchungen unter 'Nicht zugeordnet'."),
        ("T", "• Die Filialvergleiche im Dashboard und in der Jahresübersicht zeigen die ersten beiden Filialen aus 'Listen' (Zürich, Luzern). Weitere Filialen sind in 'Alle' enthalten und über die Filialauswahl einzeln auswertbar."),
        ("", ""),
        ("H", "Wichtige Hinweise"),
        ("T", "• Sicherung: Datei in OneDrive, Google Drive oder Dropbox ablegen (Versionsverlauf) und wöchentlich eine Kopie behalten."),
        ("T", "• Diese Datei ist ein Führungsinstrument (Cockpit), keine Buchhaltung nach OR. Sie ersetzt weder MWST-Abrechnung noch Jahresabschluss – liefert dem Treuhänder aber sauber kategorisierte Daten."),
        ("T", "• Die Arbeitgeberbeiträge im Blatt 'Mitarbeiter' sind Schätzwerte (AHV/IV/EO 5.3 %, ALV 1.1 %; BVG, UVG/KTG und FAK je nach Vertrag und Kanton). Bitte mit Treuhänder oder Lohnabrechnung abgleichen."),
        ("T", "• Beispieldaten (Jahr 2025, Bemerkung 'Beispiel', Mitarbeiter 'Beispiel: …', Budget-Vorbelegung) sind frei erfunden und dienen nur zur Demonstration."),
        ("T", "• Kapazität: 5'000 Buchungen. Danach entweder neue Datei pro Jahr anlegen oder den Bereich in den Formeln (Zeile 5000) erweitern."),
    ]
    r = 4
    for kind, text in lines:
        if kind == "H":
            put(ws, f"B{r}", text, bold=True, size=12, color=C_HDR)
        elif kind.startswith("L"):
            if kind == "L1":
                put(ws, f"A{r}", "Jahr", bold=True, color="0000FF", fill_=C_IN, align="center", border=BORDER)
            elif kind == "L2":
                put(ws, f"A{r}", "1'234", color="0000FF", align="center", border=BORDER)
            else:
                put(ws, f"A{r}", "!", bold=True, fill_=C_WARN, align="center", border=BORDER)
            put(ws, f"B{r}", text, wrap=True)
        elif kind == "T":
            put(ws, f"B{r}", text, wrap=True)
        if kind in ("T", "L1", "L2", "L3"):
            ws.row_dimensions[r].height = 15 * max(1, -(-len(text) // 120))
        r += 1
    widths(ws, {"A": 8, "B": 130})
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "7F7F7F"
    ws.sheet_view.showGridLines = False
    return ws


def build(path, test=False):
    wb = Workbook()
    wb.remove(wb.active)
    rows = example_rows()
    build_listen(wb)
    build_buchungen(wb, rows)
    build_monat(wb, "Luzern" if test else "Alle")
    build_jahr(wb, "Zürich" if test else "Alle")
    build_mitarbeiter(wb)
    build_budget(wb)
    build_dashboard(wb)
    build_anleitung(wb)
    order = ["Anleitung", "Dashboard", "Buchungen", "Monatsübersicht", "Jahresübersicht", "Mitarbeiter", "Budget", "Listen"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 1
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Dashboard"
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "creaDIG"
    wb.properties.title = "Vegitat – Finanzübersicht Zürich & Luzern"
    wb.save(path)
    return rows


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    out = Path(args[0]) if args else Path(__file__).with_name("Vegitat_Finanzuebersicht.xlsx")
    n = len(build(out, test="--test" in sys.argv))
    print(f"geschrieben: {out}  ({n} Beispielbuchungen)")
